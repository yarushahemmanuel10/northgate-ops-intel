"""CRM workbook: no ground-truth leakage, round trip, determinism, synthetic marker."""
import datetime as dt
import re
import time

import pytest
from openpyxl import load_workbook

from generator.config import load_config
from generator.crm import ACCOUNT_COLUMNS, GROUND_TRUTH_ONLY, LINE_COLUMNS, build
from generator.customers import SYNTHETIC_HEADER, generate, make_rngs


@pytest.fixture(scope="module")
def cfg():
    return load_config()


@pytest.fixture(scope="module")
def truth(cfg):
    return generate(cfg, *make_rngs(cfg["seed"]))


@pytest.fixture(scope="module")
def workbook(cfg, tmp_path_factory):
    path = build(cfg, tmp_path_factory.mktemp("crm"))
    return load_workbook(path)


def _rows(ws):
    """Header (row 2) and data rows (row 3+) as dicts."""
    values = list(ws.iter_rows(values_only=True))
    header = list(values[1])
    return header, [dict(zip(header, r)) for r in values[2:]]


def _normalise(header: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", header.lower()).strip("_")


def _as_date(v):
    return v.date() if isinstance(v, dt.datetime) else v


def test_sheets(workbook):
    assert workbook.sheetnames == ["README", "Accounts", "Contract Lines"]
    assert workbook["README"]["A1"].value == "SYNTHETIC DATA"


@pytest.mark.parametrize("sheet", ["Accounts", "Contract Lines"])
def test_data_sheets_start_with_synthetic_marker(workbook, sheet):
    assert workbook[sheet]["A1"].value == SYNTHETIC_HEADER


@pytest.mark.parametrize("sheet, columns", [("Accounts", ACCOUNT_COLUMNS),
                                            ("Contract Lines", LINE_COLUMNS)])
def test_headers_are_exactly_the_crm_columns(workbook, sheet, columns):
    header, _ = _rows(workbook[sheet])
    assert header == list(columns.values())


def test_no_ground_truth_columns_anywhere(workbook):
    banned = set(GROUND_TRUTH_ONLY) | {"at_risk", "is_new_in_window", "lookalike_group",
                                       "revenue_source", "lookalike_kind"}
    for ws in workbook.worksheets:
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert _normalise(value) not in banned, f"{ws.title}: {value}"


def test_accounts_round_trip(workbook, truth):
    customers, _ = truth
    _, rows = _rows(workbook["Accounts"])
    assert len(rows) == len(customers)
    by_id = {r["Account ID"]: r for r in rows}
    for c in customers.itertuples():
        r = by_id[c.account_id]
        assert r["Company Name"] == c.company_name
        assert r["Sector"] == c.sector
        assert r["Account Manager"] == c.account_manager
        assert _as_date(r["Customer Since"]) == c.customer_since
        assert r["Status"] == c.status
        assert _as_date(r["Closed Date"]) == c.closed_date


def test_closed_customers_keep_status_and_date(workbook, truth):
    customers, _ = truth
    _, rows = _rows(workbook["Accounts"])
    closed = [r for r in rows if r["Status"] == "closed"]
    assert len(closed) == (customers.status == "closed").sum()
    assert all(isinstance(_as_date(r["Closed Date"]), dt.date) for r in closed)


def test_contract_lines_round_trip(workbook, truth):
    _, lines = truth
    _, rows = _rows(workbook["Contract Lines"])
    assert len(rows) == len(lines)
    by_id = {r["Line ID"]: r for r in rows}
    for ln in lines.itertuples():
        r = by_id[ln.line_id]
        assert r["Account ID"] == ln.account_id
        assert r["Business Line"] == ln.business_line
        assert r["Annual Value (GBP)"] == ln.annual_value
        assert _as_date(r["Contract Start"]) == ln.contract_start
        assert r["Term (Years)"] == ln.term_years
        assert _as_date(r["Renewal Date"]) == ln.renewal_date
        assert _as_date(r["End Date"]) == ln.end_date


def test_shop_only_customers_are_listed_without_lines(workbook, truth):
    customers, _ = truth
    shop_only = set(customers.loc[customers.business_lines == "BL5", "account_id"])
    assert shop_only, "expected some shop-only customers in the ground truth"
    _, accounts = _rows(workbook["Accounts"])
    _, lines = _rows(workbook["Contract Lines"])
    assert shop_only <= {r["Account ID"] for r in accounts}
    assert not shop_only & {r["Account ID"] for r in lines}


def test_same_seed_gives_identical_bytes(cfg, tmp_path):
    a = build(cfg, tmp_path / "a").read_bytes()
    time.sleep(1.1)                      # make sure a wall-clock timestamp would differ
    b = build(cfg, tmp_path / "b").read_bytes()
    assert a == b


def test_different_seed_gives_different_bytes(cfg, tmp_path):
    a = build(cfg, tmp_path / "a").read_bytes()
    b = build(cfg, tmp_path / "b", seed=cfg["seed"] + 1).read_bytes()
    assert a != b
