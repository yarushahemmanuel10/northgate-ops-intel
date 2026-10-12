"""Connectors for file sources: CRM workbook (Excel), helpdesk and telecoms (CSV)."""
from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

from openpyxl import load_workbook

from pipeline.ingest.common import LINEAGE, MARKER, Rejected, SchemaError, check_header


def _cell_text(v) -> str:
    """Excel cell as the text a person would see in the workbook."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def read_excel_sheet(path: Path, sheet: str, expected: list[str], system: str, ingested_at: str,
                     rejected: Rejected) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SchemaError(f"{path.name}: sheet '{sheet}' not found (found {wb.sheetnames})")
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    if not rows or not str(rows[0][0] or "").startswith(MARKER):
        raise SchemaError(f"{path.name}/{sheet}: row 1 must be the '{MARKER}' marker")
    header = [str(c) for c in rows[1] if c is not None]
    check_header(header, expected, f"{path.name}/{sheet}")
    out = []
    for n, r in enumerate(rows[2:], start=3):                 # Excel row numbers
        values = list(r)
        if len(values) > len(expected) and any(v not in (None, "") for v in values[len(expected):]):
            rejected.add(system, f"{path.name}/{sheet}", str(n), "more values than columns", repr(values))
            continue
        if all(v in (None, "") for v in values):
            continue
        rec = {c: _cell_text(values[i] if i < len(values) else None) for i, c in enumerate(expected)}
        rec.update(dict(zip(LINEAGE, (system, f"{path.name}/{sheet}", str(n), ingested_at))))
        out.append(rec)
    return out


def read_csv(path: Path, expected: list[str], system: str, ingested_at: str, rejected: Rejected) -> list[dict]:
    with open(path, encoding="utf-8", newline="") as f:
        first = f.readline()
        if not first.startswith(MARKER):
            raise SchemaError(f"{path.name}: line 1 must be the '{MARKER}' marker")
        reader = csv.reader(f)
        try:
            header = next(reader)
        except StopIteration:
            raise SchemaError(f"{path.name}: no header row") from None
        check_header(header, expected, path.name)
        out = []
        for values in reader:
            line = reader.line_num + 1                      # +1 for the marker line read separately
            if len(values) != len(expected):
                rejected.add(system, path.name, str(line),
                             f"{len(values)} values for {len(expected)} columns", ",".join(values))
                continue
            rec = dict(zip(expected, values))
            rec.update(dict(zip(LINEAGE, (system, path.name, str(line), ingested_at))))
            out.append(rec)
    return out
